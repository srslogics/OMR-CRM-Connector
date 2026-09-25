from pathlib import Path
import os, json, sqlite3, uuid, secrets, hashlib, hmac, time, threading, io, csv
from contextlib import asynccontextmanager, contextmanager
from urllib.parse import urlparse
from fastapi import FastAPI, Request, HTTPException, UploadFile, File, Form
from fastapi.responses import FileResponse, JSONResponse, Response
from fastapi.staticfiles import StaticFiles
import cv2
import pymupdf as fitz
from processor import raster,align,detect_question,ocr_crop,score,validate_config,FIELDS,ocr_available

ROOT=Path(__file__).parent
DATA=Path(os.environ.get('PAPERDESK_DATA',ROOT/'data')).resolve();DATA.mkdir(parents=True,exist_ok=True)
DB=DATA/'paperdesk.sqlite';wake=threading.Event();stop=threading.Event();worker=None

@contextmanager
def db():
    c=sqlite3.connect(DB,timeout=30);c.row_factory=sqlite3.Row;c.execute('PRAGMA foreign_keys=ON')
    try:
        with c:yield c
    finally:c.close()

def init():
    with db() as c:
        c.executescript('''PRAGMA journal_mode=WAL;
        CREATE TABLE IF NOT EXISTS users(id INTEGER PRIMARY KEY,name TEXT,email TEXT UNIQUE,password TEXT,salt TEXT);
        CREATE TABLE IF NOT EXISTS sessions(token TEXT PRIMARY KEY,user_id INTEGER,expires REAL);
        CREATE TABLE IF NOT EXISTS exams(id TEXT PRIMARY KEY,name TEXT,class_name TEXT,config TEXT,locked INTEGER DEFAULT 0,created REAL);
        CREATE TABLE IF NOT EXISTS batches(id TEXT PRIMARY KEY,exam_id TEXT,name TEXT,pages INTEGER,total INTEGER,done INTEGER DEFAULT 0,status TEXT,error TEXT,created REAL);
        CREATE TABLE IF NOT EXISTS papers(id TEXT PRIMARY KEY,batch_id TEXT,idx INTEGER,data TEXT,status TEXT,version INTEGER DEFAULT 1,UNIQUE(batch_id,idx));
        CREATE TABLE IF NOT EXISTS audit(id INTEGER PRIMARY KEY,paper_id TEXT,user_id INTEGER,at REAL,before_data TEXT,after_data TEXT);
        ''')
        c.execute("UPDATE batches SET status='queued' WHERE status='processing'")

@asynccontextmanager
async def lifespan(app):
    global worker
    init();stop.clear();worker=threading.Thread(target=work_loop,daemon=True);worker.start();yield;stop.set();wake.set();worker.join(timeout=3)
app=FastAPI(title='PaperDesk',lifespan=lifespan,docs_url=None,redoc_url=None,openapi_url=None)
app.mount('/static',StaticFiles(directory=ROOT/'static'),name='static')

@app.middleware('http')
async def secure(request,call_next):
    if request.method not in ('GET','HEAD','OPTIONS'):
        origin=request.headers.get('origin')
        if origin and urlparse(origin).netloc!=request.headers.get('host'):return JSONResponse({'detail':'Cross-origin request rejected'},403)
    r=await call_next(request);r.headers['X-Content-Type-Options']='nosniff';r.headers['X-Frame-Options']='DENY';r.headers['Referrer-Policy']='same-origin'
    if request.url.path.startswith(('/api','/media')):r.headers['Cache-Control']='no-store'
    return r

def user(request):
    token=hashlib.sha256(request.cookies.get('paperdesk','').encode()).hexdigest()
    with db() as c:r=c.execute('SELECT u.id,u.name,u.email FROM sessions s JOIN users u ON u.id=s.user_id WHERE token=? AND expires>?',(token,time.time())).fetchone()
    if not r:raise HTTPException(401,'Please sign in.')
    return dict(r)

def require_row(table,id):
    with db() as c:r=c.execute(f'SELECT * FROM {table} WHERE id=?',(id,)).fetchone()
    if not r:raise HTTPException(404,'Record not found')
    return dict(r)

def issue_session(uid):
    token=secrets.token_urlsafe(32)
    with db() as c:
        c.execute('DELETE FROM sessions WHERE expires<?',(time.time(),));c.execute('INSERT INTO sessions VALUES(?,?,?)',(hashlib.sha256(token.encode()).hexdigest(),uid,time.time()+43200))
    r=JSONResponse({'ok':True});r.set_cookie('paperdesk',token,httponly=True,samesite='strict',secure=os.getenv('COOKIE_SECURE')=='1',max_age=43200);return r

def password_hash(p,s):return hashlib.pbkdf2_hmac('sha256',p.encode(),bytes.fromhex(s),260000).hex()
login_attempts={}
@app.get('/api/session')
def session(request:Request):
    with db() as c:setup=c.execute('SELECT count(*) FROM users').fetchone()[0]==0
    try:u=user(request)
    except HTTPException:u=None
    return {'setup_required':setup,'user':u,'ocr_available':ocr_available()}

@app.post('/api/setup')
async def setup(request:Request):
    p=await request.json();name=str(p.get('name','')).strip();email=str(p.get('email','')).strip().lower();password=str(p.get('password',''))
    if not name or '@' not in email or len(password)<10:raise HTTPException(400,'Name, email and a password of at least 10 characters are required.')
    with db() as c:
        c.execute('BEGIN IMMEDIATE')
        if c.execute('SELECT count(*) FROM users').fetchone()[0]:raise HTTPException(409,'Administrator is already configured.')
        salt=secrets.token_hex(16);cur=c.execute('INSERT INTO users(name,email,password,salt) VALUES(?,?,?,?)',(name[:80],email[:200],password_hash(password,salt),salt));uid=cur.lastrowid
    return issue_session(uid)

@app.post('/api/login')
async def login(request:Request):
    ip=request.client.host;recent=[t for t in login_attempts.get(ip,[]) if t>time.time()-60]
    if len(recent)>=8:raise HTTPException(429,'Too many attempts. Try again in a minute.')
    login_attempts[ip]=recent+[time.time()];p=await request.json()
    with db() as c:u=c.execute('SELECT * FROM users WHERE email=?',(str(p.get('email','')).strip().lower(),)).fetchone()
    if not u or not hmac.compare_digest(password_hash(str(p.get('password','')),u['salt']),u['password']):raise HTTPException(401,'Email or password is incorrect.')
    login_attempts.pop(ip,None);return issue_session(u['id'])

@app.post('/api/logout')
def logout(request:Request):
    token=hashlib.sha256(request.cookies.get('paperdesk','').encode()).hexdigest()
    with db() as c:c.execute('DELETE FROM sessions WHERE token=?',(token,))
    r=JSONResponse({'ok':True});r.delete_cookie('paperdesk');return r

@app.get('/api/dashboard')
def dashboard(request:Request):
    user(request)
    with db() as c:
        exams=[dict(r) for r in c.execute('SELECT id,name,class_name,locked,created FROM exams ORDER BY created DESC')]
        batches=[dict(r) for r in c.execute('SELECT b.*,e.name exam_name,e.class_name,(SELECT count(*) FROM papers p WHERE p.batch_id=b.id AND p.status="approved") approved FROM batches b JOIN exams e ON e.id=b.exam_id ORDER BY b.created DESC')]
    return {'exams':exams,'batches':batches}

@app.post('/api/exams')
async def create_exam(request:Request):
    user(request);p=await request.json();name=str(p.get('name','')).strip();cl=str(p.get('class_name','')).strip()
    if not name or not cl:raise HTTPException(400,'Exam name and class are required.')
    id=uuid.uuid4().hex
    with db() as c:c.execute('INSERT INTO exams VALUES(?,?,?,?,?,?)',(id,name[:120],cl[:40],json.dumps({'key':['']*25,'mapping':{},'fields':{}}),0,time.time()))
    (DATA/'exams'/id).mkdir(parents=True);return {'id':id}

@app.get('/api/exams/{id}')
def exam(id:str,request:Request):
    user(request);e=require_row('exams',id);e['config']=json.loads(e['config']);e['has_template']=(DATA/'exams'/id/'page-0.png').exists();return e

@app.put('/api/exams/{id}')
async def save_exam(id:str,request:Request):
    user(request);e=require_row('exams',id)
    if e['locked']:raise HTTPException(409,'This exam is locked. Create a new exam for a different key or template.')
    conf=await request.json()
    if len(json.dumps(conf))>80000:raise HTTPException(400,'Configuration too large')
    if not isinstance(conf.get('mapping'),dict) or not isinstance(conf.get('fields'),dict) or not isinstance(conf.get('key'),list):raise HTTPException(400,'Invalid configuration')
    with db() as c:c.execute('UPDATE exams SET config=? WHERE id=?',(json.dumps(conf),id))
    return {'ok':True}

async def save_upload(file,target,limit=150*1024*1024):
    size=0
    try:
        with open(target,'wb') as f:
            while chunk:=await file.read(1024*1024):
                size+=len(chunk)
                if size>limit:raise HTTPException(413,'PDF exceeds the 150 MB upload limit. Compress the scan or use smaller batches.')
                f.write(chunk)
        with fitz.open(target) as doc:
            if not doc.is_pdf or doc.needs_pass:raise ValueError('Use an unencrypted PDF.')
            pages=len(doc)
        return pages
    except Exception as exc:
        Path(target).unlink(missing_ok=True)
        if isinstance(exc,HTTPException):raise
        raise HTTPException(400,'Cannot read this PDF. Use a valid, unencrypted PDF.')

@app.post('/api/exams/{id}/template')
async def upload_template(id:str,request:Request,file:UploadFile=File(...)):
    user(request);e=require_row('exams',id)
    if e['locked']:raise HTTPException(409,'Exam is locked.')
    folder=DATA/'exams'/id;folder.mkdir(parents=True,exist_ok=True);target=folder/'incoming.pdf';n=await save_upload(file,target)
    if n!=2:target.unlink();raise HTTPException(400,'The blank master must have exactly two pages.')
    try:
        with fitz.open(target) as doc:
            for i in range(2):cv2.imwrite(str(folder/f'page-{i}.png'),raster(doc[i]))
        target.replace(folder/'template.pdf')
    except Exception:raise HTTPException(400,'Unable to render the master PDF.')
    conf=json.loads(e['config']);conf['mapping']={};conf['fields']={}
    with db() as c:c.execute('UPDATE exams SET config=? WHERE id=?',(json.dumps(conf),id))
    return {'ok':True}

@app.post('/api/exams/{id}/lock')
def lock(id:str,request:Request):
    user(request);e=require_row('exams',id)
    if not (DATA/'exams'/id/'page-1.png').exists():raise HTTPException(400,'Upload a two-page blank master first.')
    try:validate_config(json.loads(e['config']))
    except ValueError as err:raise HTTPException(400,str(err))
    with db() as c:c.execute('UPDATE exams SET locked=1 WHERE id=?',(id,))
    return {'ok':True}

@app.post('/api/batches')
async def upload_batch(request:Request,exam_id:str=Form(...),file:UploadFile=File(...)):
    user(request);e=require_row('exams',exam_id)
    if not e['locked']:raise HTTPException(400,'Verify and lock the exam first.')
    id=uuid.uuid4().hex;folder=DATA/'batches'/id;folder.mkdir(parents=True);n=await save_upload(file,folder/'source.pdf')
    if n<2 or n>600 or n%2:
        (folder/'source.pdf').unlink();raise HTTPException(400,'Use an even page count between 2 and 600. Each student must have two consecutive pages.')
    with db() as c:c.execute('INSERT INTO batches VALUES(?,?,?,?,?,0,?,?,?)',(id,exam_id,Path(file.filename or 'Batch.pdf').name[:150],n,n//2,'queued','',time.time()))
    wake.set();return {'id':id}

@app.get('/api/batches/{id}')
def batch(id:str,request:Request):
    user(request);b=require_row('batches',id)
    with db() as c:rows=c.execute('SELECT * FROM papers WHERE batch_id=? ORDER BY idx',(id,)).fetchall()
    b['papers']=[]
    for row in rows:
        p=dict(row);p['data']=json.loads(p['data']);b['papers'].append(p)
    return b

@app.post('/api/batches/{id}/retry')
def retry(id:str,request:Request):
    user(request);b=require_row('batches',id)
    if b['status']!='failed':raise HTTPException(409,'Only failed batches can be retried.')
    with db() as c:c.execute("UPDATE batches SET status='queued',error='' WHERE id=?",(id,))
    wake.set();return {'ok':True}

@app.get('/api/papers/{id}')
def paper(id:str,request:Request):
    user(request);p=require_row('papers',id);p['data']=json.loads(p['data']);b=require_row('batches',p['batch_id']);e=require_row('exams',b['exam_id']);p['exam']=dict(e,config=json.loads(e['config']));return p

@app.put('/api/papers/{id}')
async def review(id:str,request:Request):
    u=user(request);p=require_row('papers',id);incoming=await request.json();old=json.loads(p['data']);answers=incoming.get('answers');fields=incoming.get('fields',{})
    if not isinstance(answers,list) or len(answers)!=25 or any(a not in ['A','B','C','D','-','?'] for a in answers):raise HTTPException(400,'Provide 25 valid answers.')
    if not isinstance(fields,dict):raise HTTPException(400,'Invalid fields')
    fields={k:str(fields.get(k,'')).strip()[:180] for k in FIELDS}
    approve=incoming.get('approve') is True
    if approve and ('?' in answers or not fields['name'] or incoming.get('pairing_verified') is not True):raise HTTPException(400,'Resolve every answer, enter the student name and verify page pairing before approval.')
    e=require_row('exams',require_row('batches',p['batch_id'])['exam_id']);new={**old,'answers':answers,'fields':fields,'score':score(answers,json.loads(e['config'])['key']),'pairing_verified':bool(incoming.get('pairing_verified'))}
    with db() as c:
        cur=c.execute('UPDATE papers SET data=?,status=?,version=version+1 WHERE id=? AND version=?',(json.dumps(new),'approved' if approve else 'review',id,incoming.get('version')))
        if cur.rowcount!=1:raise HTTPException(409,'This paper changed in another window. Reload before saving.')
        c.execute('INSERT INTO audit(paper_id,user_id,at,before_data,after_data) VALUES(?,?,?,?,?)',(id,u['id'],time.time(),p['data'],json.dumps(new)))
    return {'ok':True}

@app.get('/media/exams/{id}/{page}.png')
def master_image(id:str,page:int,request:Request):
    user(request);require_row('exams',id)
    if page not in (0,1):raise HTTPException(404)
    p=DATA/'exams'/id/f'page-{page}.png'
    if not p.exists():raise HTTPException(404)
    return FileResponse(p)

@app.get('/media/papers/{id}/{page}.png')
def paper_image(id:str,page:int,request:Request):
    user(request);p=require_row('papers',id)
    if page not in (0,1):raise HTTPException(404)
    return FileResponse(DATA/'batches'/p['batch_id']/f"{p['idx']}-{page}.png")

def work_loop():
    while not stop.is_set():
        with db() as c:b=c.execute("SELECT * FROM batches WHERE status='queued' ORDER BY created LIMIT 1").fetchone()
        if not b:wake.wait(1);wake.clear();continue
        try:process_batch(dict(b))
        except Exception as e:
            with db() as c:c.execute("UPDATE batches SET status='failed',error=? WHERE id=?",('Processing failed: '+str(e)[:180],b['id']))

def process_batch(b):
    e=require_row('exams',b['exam_id']);conf=json.loads(e['config']);folder=DATA/'batches'/b['id'];masters=[cv2.imread(str(DATA/'exams'/e['id']/f'page-{i}.png')) for i in range(2)]
    with db() as c:c.execute("UPDATE batches SET status='processing',error='' WHERE id=?",(b['id'],))
    with fitz.open(folder/'source.pdf') as doc:
        for idx in range(b['total']):
            if stop.is_set():return
            with db() as c:exists=c.execute('SELECT 1 FROM papers WHERE batch_id=? AND idx=?',(b['id'],idx)).fetchone()
            if exists:continue
            fields={k:'' for k in FIELDS};fields['class']=e['class_name'];answers=['?']*25;details=[{'reason':'Page alignment needs review','evidence':[]} for _ in range(25)];flags=[]
            for pageno in range(2):
                im=raster(doc[idx*2+pageno]);aligned,quality=align(im,masters[pageno])
                if aligned is None:
                    flags.append(f'Page {pageno+1}: alignment failed; check page order and scan quality.')
                    cv2.imwrite(str(folder/f'{idx}-{pageno}.png'),im);continue
                cv2.imwrite(str(folder/f'{idx}-{pageno}.png'),aligned)
                for q,m in conf['mapping'].items():
                    if m['page']!=pageno:continue
                    d=detect_question(aligned,masters[pageno],m['boxes']);answers[int(q)-1]=d['answer'];details[int(q)-1]=d
                for k,m in conf.get('fields',{}).items():
                    if m['page']==pageno:fields[k]=ocr_crop(aligned,m['box'],folder,k)
            flags.append('Verify student details, both pages and all detected answers before approving.')
            data={'fields':fields,'answers':answers,'details':details,'flags':flags,'pairing_verified':False,'score':score(answers,conf['key'])}
            with db() as c:
                c.execute('INSERT INTO papers VALUES(?,?,?,?,?,1)',(uuid.uuid4().hex,b['id'],idx,json.dumps(data),'review'))
                c.execute('UPDATE batches SET done=(SELECT count(*) FROM papers WHERE batch_id=?) WHERE id=?',(b['id'],b['id']))
    with db() as c:c.execute("UPDATE batches SET status='ready' WHERE id=?",(b['id'],))

@app.get('/api/batches/{id}/export/{kind}')
def export(id:str,kind:str,request:Request):
    user(request);b=require_row('batches',id)
    with db() as c:rows=c.execute('SELECT data,idx FROM papers WHERE batch_id=? AND status="approved" ORDER BY idx',(id,)).fetchall()
    if not rows:raise HTTPException(400,'Approve at least one paper before exporting.')
    headers=['Student number']+FIELDS+['Science /40','Mathematics /40','Mental ability /20','Total /100','Correct','Blank']+[f'Q{i}' for i in range(1,26)]
    values=[]
    for r in rows:
        d=json.loads(r['data']);s=d['score'];values.append([r['idx']+1]+[d['fields'].get(k,'') for k in FIELDS]+[s['science'],s['mathematics'],s['mental_ability'],s['total'],s['correct'],s['blank']]+d['answers'])
    def safe(v):return "'"+v if isinstance(v,str) and v.startswith(('=','+','-','@','\t','\r')) else v
    values=[[safe(v) for v in r] for r in values]
    if kind=='csv':
        out=io.StringIO();w=csv.writer(out);w.writerow(headers);w.writerows(values);content=out.getvalue().encode('utf-8-sig');mime='text/csv'
    elif kind=='xlsx':
        from openpyxl import Workbook
        from openpyxl.styles import Font,PatternFill
        wb=Workbook();ws=wb.active;ws.title='Approved results';ws.append(headers)
        for r in values:ws.append(r)
        for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='162C42')
        ws.freeze_panes='B2';ws.auto_filter.ref=ws.dimensions
        from openpyxl.utils import get_column_letter
        for i in range(1,len(headers)+1):ws.column_dimensions[get_column_letter(i)].width=22 if i<10 else 14
        out=io.BytesIO();wb.save(out);content=out.getvalue();mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    else:raise HTTPException(404)
    return Response(content,media_type=mime,headers={'Content-Disposition':f'attachment; filename="approved-results.{kind}"'})

@app.get('/api/papers/{id}/marksheet')
def marksheet(id:str,request:Request):
    user(request);p=require_row('papers',id)
    if p['status']!='approved':raise HTTPException(400,'Approve this paper first.')
    from reportlab.pdfgen import canvas
    from reportlab.lib.colors import HexColor
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    fontpath=ROOT/'assets'/'NotoSans-Regular.ttf'
    font='Helvetica'
    if fontpath.exists():
        if 'ReportFont' not in pdfmetrics.getRegisteredFontNames():pdfmetrics.registerFont(TTFont('ReportFont',str(fontpath)))
        font='ReportFont'
    d=json.loads(p['data']);b=require_row('batches',p['batch_id']);e=require_row('exams',b['exam_id']);out=io.BytesIO();c=canvas.Canvas(out,pagesize=(595,842));c.setTitle('Student marksheet')
    c.setFillColor(HexColor('#162c42'));c.rect(0,725,595,117,fill=1,stroke=0);c.setFillColor(HexColor('#ffffff'));c.setFont(font,25);c.drawString(42,784,'Student marksheet');c.setFont(font,12);c.drawString(42,755,e['name'][:65]);c.setFillColor(HexColor('#172b43'));c.setFont(font,12);y=693
    for k in ['name','school','class','section','taluka','district']:
        c.drawString(42,y,f"{k.replace('_',' ').title()}: {d['fields'].get(k,'')[:60]}");y-=25
    y-=18
    for key,label in [('science','Science /40'),('mathematics','Mathematics /40'),('mental_ability','Mental ability /20'),('total','Total /100')]:
        c.drawString(42,y,label);c.drawRightString(545,y,str(d['score'][key]));y-=27
    c.setFont(font,10);c.drawString(42,y-15,'Question-wise responses');y-=44
    key=json.loads(e['config'])['key']
    for i,a in enumerate(d['answers']):
        col=i//13;row=i%13;c.drawString(42+col*270,y-row*16,f"Q{i+1:02}   {a}    {'Correct' if a==key[i] else 'Blank' if a=='-' else 'Incorrect'}")
    c.setFont(font,9);c.drawString(42,40,'Reviewed result | PaperDesk by SrS Logics');c.save()
    return Response(out.getvalue(),media_type='application/pdf',headers={'Content-Disposition':'attachment; filename="marksheet.pdf"'})

@app.get('/')
def home():return FileResponse(ROOT/'static'/'index.html')

@app.post('/api/sample')
def sample(request:Request):
    user(request)
    from sample import create_pdf
    eid=uuid.uuid4().hex;ef=DATA/'exams'/eid;ef.mkdir(parents=True)
    conf=create_pdf(ef/'template.pdf');validate_config(conf)
    with fitz.open(ef/'template.pdf') as doc:
        for i in range(2):cv2.imwrite(str(ef/f'page-{i}.png'),raster(doc[i]))
    bid=uuid.uuid4().hex;bf=DATA/'batches'/bid;bf.mkdir(parents=True)
    create_pdf(bf/'source.pdf',['Sample Student One','Sample Student Two','Sample Student Three'])
    with db() as c:
        c.execute('INSERT INTO exams VALUES(?,?,?,?,1,?)',(eid,'Synthetic walkthrough - NOT official','9',json.dumps(conf),time.time()))
        c.execute('INSERT INTO batches VALUES(?,?,?,?,?,0,?,?,?)',(bid,eid,'Synthetic sample - 3 students.pdf',6,3,'queued','',time.time()))
    wake.set();return {'id':bid}

@app.get('/api/batches/{id}/source')
def source_pdf(id:str,request:Request):
    user(request);require_row('batches',id)
    return FileResponse(DATA/'batches'/id/'source.pdf',media_type='application/pdf',filename='original-student-batch.pdf')
